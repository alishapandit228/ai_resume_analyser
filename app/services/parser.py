"""
Resume parsing service.

Responsibilities:
  1. Extract raw text from an uploaded PDF or DOCX file.
  2. Parse that raw text into a structured dict: contact info, education,
     categorized skills, experience, projects, certifications, achievements.

Everything here is rule-based (regex + keyword matching against the skills
taxonomy) so it works with zero external API calls. If an LLM provider is
configured (see services/llm_client.py), routers can additionally ask the
LLM to refine/augment this structured output — but the app must never
*invent* information that isn't in the resume, so the LLM step is used only
for light cleanup, not fabrication.
"""
import json
import re
from pathlib import Path
from typing import Dict, List, Any

import docx  # python-docx
from pypdf import PdfReader

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
with open(DATA_DIR / "skills_taxonomy.json") as f:
    SKILLS_TAXONOMY: Dict[str, List[str]] = json.load(f)

ALL_SKILLS = {s.lower(): cat for cat, skills in SKILLS_TAXONOMY.items() for s in skills}

EMAIL_RE = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s.-]?)?(?:\(?\d{2,4}\)?[\s.-]?){2,4}\d{3,4}")
LINKEDIN_RE = re.compile(r"(https?://)?(www\.)?linkedin\.com/in/[A-Za-z0-9\-_/]+", re.I)
GITHUB_RE = re.compile(r"(https?://)?(www\.)?github\.com/[A-Za-z0-9\-_/]+", re.I)
CGPA_RE = re.compile(r"(?:cgpa|gpa)\s*[:\-]?\s*(\d\.\d{1,2})", re.I)
PERCENT_RE = re.compile(r"(\d{2}(?:\.\d{1,2})?)\s*%")
YEAR_RE = re.compile(r"(20\d{2}|19\d{2})")

SECTION_HEADERS = {
    "education": ["education", "academic background", "qualification"],
    "skills": ["skills", "technical skills", "core competencies"],
    "experience": ["experience", "work experience", "employment history", "internship"],
    "projects": ["projects", "academic projects", "personal projects"],
    "certifications": ["certifications", "certificates", "licenses"],
    "achievements": ["achievements", "awards", "honors"],
}


class ParsingError(Exception):
    pass


def extract_text(file_path: str, filename: str) -> str:
    """Extract raw text from a PDF or DOCX file."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            document = docx.Document(file_path)
            text = "\n".join(p.text for p in document.paragraphs)
            for table in document.tables:
                for row in table.rows:
                    text += "\n" + " | ".join(cell.text for cell in row.cells)
        else:
            raise ParsingError(f"Unsupported file type: {ext}")
    except ParsingError:
        raise
    except Exception as e:
        raise ParsingError(f"Could not read {ext} file: {e}")

    text = text.strip()
    if not text or len(text) < 30:
        raise ParsingError(
            "The resume appears to be empty or the text could not be read "
            "(possibly a scanned/image-based PDF). Please upload a text-based "
            "PDF or DOCX file."
        )
    return text


def _split_into_sections(text: str) -> Dict[str, str]:
    """Naively splits resume text into sections based on common headings."""
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {"header": []}
    current = "header"

    for line in lines:
        stripped = line.strip().lower().strip(":")
        matched_section = None
        for key, headers in SECTION_HEADERS.items():
            if stripped in headers or (len(stripped) < 40 and any(h == stripped for h in headers)):
                matched_section = key
                break
        if matched_section:
            current = matched_section
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items()}


def _extract_contact(text: str, header_text: str) -> Dict[str, Any]:
    email = EMAIL_RE.search(text)
    linkedin = LINKEDIN_RE.search(text)
    github = GITHUB_RE.search(text)
    phone = PHONE_RE.search(header_text) or PHONE_RE.search(text[:500])

    # Best-effort name guess: first non-empty line that isn't an email/phone/link
    name = None
    for line in header_text.split("\n"):
        candidate = line.strip()
        if not candidate:
            continue
        if EMAIL_RE.search(candidate) or PHONE_RE.search(candidate) or "http" in candidate.lower():
            continue
        if 2 <= len(candidate.split()) <= 5 and len(candidate) < 60:
            name = candidate
            break

    return {
        "name": name,
        "email": email.group(0) if email else None,
        "phone": phone.group(0).strip() if phone else None,
        "linkedin": linkedin.group(0) if linkedin else None,
        "github": github.group(0) if github else None,
        "portfolio": None,
    }


def _extract_education(edu_text: str) -> List[Dict[str, Any]]:
    if not edu_text:
        return []
    entries = []
    blocks = [b.strip() for b in re.split(r"\n{1,2}", edu_text) if b.strip()]
    for block in blocks:
        cgpa_match = CGPA_RE.search(block)
        percent_match = PERCENT_RE.search(block)
        year_match = YEAR_RE.findall(block)
        entries.append({
            "raw": block,
            "degree": block.split("\n")[0][:120],
            "cgpa_or_percentage": (cgpa_match.group(1) if cgpa_match
                                    else (f"{percent_match.group(1)}%" if percent_match else None)),
            "graduation_year": year_match[-1] if year_match else None,
        })
    return entries


def _extract_skills(full_text: str) -> Dict[str, List[str]]:
    lower = full_text.lower()
    found: Dict[str, List[str]] = {cat: [] for cat in SKILLS_TAXONOMY}
    for skill_lower, category in ALL_SKILLS.items():
        # word-boundary-ish match, tolerant of punctuation like "c++"
        pattern = re.escape(skill_lower)
        if re.search(rf"(?<![a-z0-9]){pattern}(?![a-z0-9])", lower):
            found[category].append(skill_lower)
    return {cat: sorted(set(skills)) for cat, skills in found.items() if skills}


def _extract_experience(exp_text: str) -> List[Dict[str, Any]]:
    if not exp_text:
        return []
    blocks = [b.strip() for b in re.split(r"\n{2,}", exp_text) if b.strip()]
    if len(blocks) <= 1:
        blocks = [b.strip() for b in exp_text.split("\n") if b.strip()]
    entries = []
    for block in blocks[:8]:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        title_line = lines[0] if lines else block[:100]
        years = YEAR_RE.findall(block)
        duration = None
        if len(years) >= 2:
            duration = f"{years[0]} - {years[-1]}"
        elif years:
            duration = years[0]
        entries.append({
            "title_line": title_line,
            "duration": duration,
            "responsibilities": lines[1:] if len(lines) > 1 else [],
        })
    return entries


def _extract_projects(proj_text: str) -> List[Dict[str, Any]]:
    if not proj_text:
        return []
    blocks = [b.strip() for b in re.split(r"\n{2,}", proj_text) if b.strip()]
    if len(blocks) <= 1:
        blocks = [b.strip() for b in proj_text.split("\n") if b.strip()]
    entries = []
    for block in blocks[:10]:
        lines = [l.strip() for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        name = lines[0][:100]
        description = " ".join(lines[1:]) if len(lines) > 1 else ""
        techs = [s for s in ALL_SKILLS if re.search(rf"(?<![a-z0-9]){re.escape(s)}(?![a-z0-9])", block.lower())]
        entries.append({
            "name": name,
            "description": description[:500],
            "technologies": sorted(set(techs)),
        })
    return entries


def parse_resume(text: str) -> Dict[str, Any]:
    """Turn raw resume text into a structured dict. Extraction only — never
    fabricates data that isn't present."""
    sections = _split_into_sections(text)

    contact = _extract_contact(text, sections.get("header", ""))
    education = _extract_education(sections.get("education", ""))
    skills = _extract_skills(text)
    experience = _extract_experience(sections.get("experience", ""))
    projects = _extract_projects(sections.get("projects", ""))
    certifications = [l.strip("-• \t") for l in sections.get("certifications", "").split("\n") if l.strip()]
    achievements = [l.strip("-• \t") for l in sections.get("achievements", "").split("\n") if l.strip()]

    flat_skills = sorted({s for skills_list in skills.values() for s in skills_list})

    return {
        "contact": contact,
        "education": education,
        "skills": skills,
        "flat_skills": flat_skills,
        "experience": experience,
        "projects": projects,
        "certifications": certifications,
        "achievements": achievements,
        "word_count": len(text.split()),
    }
